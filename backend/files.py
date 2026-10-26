"""Uploaded file storage and a cheap shape inspection (rows, cols, header).

Inspection is deliberately light (no pandas) so the app stays dependency-free.
The engine re-reads the file however it likes.
"""

from __future__ import annotations

import asyncio
import codecs
import contextlib
import csv
import math
import threading
import uuid
from collections import OrderedDict
from collections.abc import Iterator
from dataclasses import dataclass, field
from itertools import islice
from pathlib import Path

from fastapi import HTTPException, UploadFile

from .config import settings

# Default is 128 KB per cell, which a single long text cell can exceed.
csv.field_size_limit(2**31 - 1)

ALLOWED_EXTENSIONS = {".csv", ".tsv", ".xlsx"}
CHUNK = 1024 * 1024
PREVIEW_ROWS = 8
PREVIEW_CELL = 200  # characters of a cell sent to the upload dialog


@dataclass
class FileInfo:
    file_id: str
    filename: str
    path: Path
    size_bytes: int
    rows: int = 0
    cols: int = 0
    columns: list[str] = field(default_factory=list)
    preview: list[list[str]] = field(default_factory=list)  # first data rows, as text
    encoding: str = "utf-8-sig"  # how a .csv/.tsv is decoded (set by inspect)
    deduped: bool = False  # some repeated column names were renamed
    warnings: list[str] = field(default_factory=list)  # import problems, shown in the upload dialog

    def public(self) -> dict:
        """Shape returned by POST /api/upload."""
        return {
            "file_id": self.file_id,
            "filename": self.filename,
            "rows": self.rows,
            "cols": self.cols,
            "size_bytes": self.size_bytes,
            "columns": self.columns,
            "preview": self.preview,
            "warnings": self.warnings,
        }


@dataclass
class Selection:
    """What the user chose in the upload dialog. Column numbers index FileInfo.columns."""

    columns: list[int] | None = None  # None = every column
    target: int | None = None
    exclude: list[tuple[int, str]] = field(
        default_factory=list
    )  # drop rows where column == value


_files: dict[str, FileInfo] = {}


def get_file(file_id: str) -> FileInfo | None:
    return _files.get(file_id)


async def save_upload(upload: UploadFile) -> FileInfo:
    filename = Path(upload.filename or "upload.csv").name
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            415, f"only {', '.join(sorted(ALLOWED_EXTENSIONS))} files, please"
        )

    settings.upload_dir.mkdir(parents=True, exist_ok=True)
    file_id = uuid.uuid4().hex[:12]
    path = settings.upload_dir / f"{file_id}{ext}"
    limit = settings.max_upload_mb * 1024 * 1024

    committed = False
    try:
        size = 0
        with path.open("wb") as out:
            while chunk := await upload.read(CHUNK):
                size += len(chunk)
                if size > limit:
                    raise HTTPException(
                        413, f"file is bigger than {settings.max_upload_mb} MB"
                    )
                out.write(chunk)

        info = FileInfo(file_id=file_id, filename=filename, path=path, size_bytes=size)
        try:
            await asyncio.to_thread(inspect, info)
        except HTTPException:
            raise
        except Exception as exc:  # corrupt / unreadable file: explain instead of a bare 500
            raise HTTPException(
                422, f"could not read that file ({type(exc).__name__}: {exc})"
            ) from exc
        _files[file_id] = info
        committed = True
        return info
    finally:
        # Any way out except success (too big, unreadable, client gone mid-upload) leaves no partial file.
        if not committed:
            with contextlib.suppress(OSError):
                path.unlink(missing_ok=True)


def inspect(info: FileInfo) -> None:
    """Fill rows / cols / columns in place, plus a list of import warnings."""
    info.warnings = []
    info.deduped = False
    if info.path.suffix.lower() != ".xlsx":
        info.encoding = _detect_encoding(info.path)
        if info.encoding not in ("utf-8-sig", "utf-16"):
            info.warnings.append(
                f"This file is not valid UTF-8, so it was read as {info.encoding}"
                " (the usual Excel encoding). If accents look wrong, save it as CSV UTF-8 and upload it again."
            )
    rows = iter_rows(info)
    header = next(rows, None)
    if not header:
        raise HTTPException(422, "that file looks empty")
    names = [str(c).strip() or f"column_{i + 1}" for i, c in enumerate(header)]
    info.columns, renamed = _unique_names(names)
    info.cols = len(info.columns)
    if renamed:
        info.deduped = True
        shown = ", ".join(f"{old!r} -> {new!r}" for old, new in renamed[:5])
        more = f" (and {len(renamed) - 5} more)" if len(renamed) > 5 else ""
        info.warnings.append(
            f"Some columns share a name, so the later ones were renamed: {shown}{more}."
        )
    info.preview = []
    info.rows = 0
    delimited = info.path.suffix.lower() != ".xlsx"
    wider: list[int] = []  # data rows with cells beyond the header
    shorter: list[int] = []  # data rows that stop before the last header column
    n_wider = n_shorter = 0
    for row in rows:
        info.rows += 1
        if info.rows <= PREVIEW_ROWS:
            info.preview.append([cell_text(v)[:PREVIEW_CELL] for v in row[: info.cols]])
        used = len(row)
        while used and not cell_text(row[used - 1]):
            used -= 1  # trailing blanks (a stray comma, Excel padding) are not extra cells
        if used > info.cols:
            n_wider += 1
            if len(wider) < 3:
                wider.append(info.rows)
        elif delimited and len(row) < info.cols:
            n_shorter += 1
            if len(shorter) < 3:
                shorter.append(info.rows)
    if info.rows == 0:
        raise HTTPException(422, "that file has a header but no data rows")
    if n_wider:
        info.warnings.append(
            f"{n_wider:,} row(s) have more cells than the header has columns (for example data row"
            f" {', '.join(map(str, wider))}). The extra cells are ignored."
        )
    if n_shorter:
        info.warnings.append(
            f"{n_shorter:,} row(s) stop before the last column (for example data row"
            f" {', '.join(map(str, shorter))}). The missing cells count as blank."
        )


def _unique_names(names: list[str]) -> tuple[list[str], list[tuple[str, str]]]:
    """Make repeated column names unique the way pandas does (amount, amount.1, amount.2).

    Returns the new names and the (old, new) pairs that changed.
    """
    taken = set(names)
    seen: set[str] = set()
    out: list[str] = []
    renamed: list[tuple[str, str]] = []
    for name in names:
        new = name
        if name in seen:
            n = 1
            while f"{name}.{n}" in taken:
                n += 1
            new = f"{name}.{n}"
            taken.add(new)
            renamed.append((name, new))
        seen.add(new)
        out.append(new)
    return out, renamed


def cell_text(value) -> str:
    """A cell as the text the user sees (and filters on)."""
    return "" if value is None else str(value).strip()


def parse_number(v: str) -> float | None:
    try:
        x = float(v)  # like pandas: "1,200" and "$3" are text
    except ValueError:
        return None
    return x if math.isfinite(x) else None


def cell_matcher(value: str):
    """Test for an exclude rule: the same text, or the same number (so 0 matches 0.0)."""
    num = parse_number(value)
    return lambda cell: cell == value or (num is not None and parse_number(cell) == num)


def preview_page(
    info: FileInfo,
    rules: list[tuple[int, str]],
    show_excluded: bool,
    cursor: int,
    limit: int,
) -> dict:
    """Up to `limit` data rows for the upload dialog, after the first `cursor` rows of the file.

    With `show_excluded` off, rows an exclude rule drops are skipped (so the page shows what will be
    analysed); otherwise they are returned and flagged. `next` is the cursor for the following page.
    """
    matchers = [(j, cell_matcher(value)) for j, value in rules]
    rows: list[list[str]] = []
    flags: list[bool] = []
    it, i = _take_pager(info, cursor)  # i = data rows consumed so far
    for raw in it:
        i += 1
        cells = [cell_text(raw[j]) if j < len(raw) else "" for j in range(info.cols)]
        hit = any(match(cells[j]) for j, match in matchers)
        if hit and not show_excluded:
            continue
        rows.append([c[:PREVIEW_CELL] for c in cells])
        flags.append(hit)
        if len(rows) >= limit:
            _park_pager(info, i, it)  # the next page carries on from here
            return {"rows": rows, "excluded": flags, "next": i, "done": False}
    return {"rows": rows, "excluded": flags, "next": i, "done": True}


# Paging through a big file must not re-read it from the top for every page. A page that stops
# part-way parks its half-read iterator under (file, next cursor); the request for the next page
# takes it back. A miss (first page, a new rule, a restart) falls back to reading from the top.
_PAGERS: OrderedDict[tuple[str, int], Iterator[list]] = OrderedDict()
_PAGERS_MAX = 4  # each one holds an open file
_pagers_lock = threading.Lock()


def _take_pager(info: FileInfo, cursor: int) -> tuple[Iterator[list], int]:
    """An iterator positioned after `cursor` data rows, and how many rows it really consumed."""
    with _pagers_lock:
        it = _PAGERS.pop((info.file_id, cursor), None)
    if it is not None:
        return it, cursor
    it = iter_rows(info)
    next(it, None)  # header
    return it, sum(1 for _ in islice(it, cursor))


def _park_pager(info: FileInfo, consumed: int, it: Iterator[list]) -> None:
    with _pagers_lock:
        _PAGERS[(info.file_id, consumed)] = it
        old = [_PAGERS.popitem(last=False)[1] for _ in range(len(_PAGERS) - _PAGERS_MAX)]
    for stale in old:
        stale.close()  # type: ignore[attr-defined]  # generators; releases the file


def iter_rows(info: FileInfo) -> Iterator[list]:
    """Yield the header, then the data rows, skipping blank lines.

    Shared by the upload check and the engine so they always agree.
    """
    if info.path.suffix.lower() == ".xlsx":
        yield from _iter_xlsx(info)
    else:
        yield from _iter_delimited(info)


def _detect_encoding(path) -> str:
    """utf-16 (Excel's "Unicode Text"), utf-8 if the whole file is valid UTF-8, else Windows-1252.

    Windows-1252 is what Excel writes as "CSV" on most Windows setups; Latin-1 is the last resort
    (it accepts every byte), so reading never produces replacement characters.
    """
    with path.open("rb") as fh:
        if fh.read(4).startswith((b"\xff\xfe", b"\xfe\xff")):
            return "utf-16"
        for name in ("utf-8-sig", "cp1252"):
            fh.seek(0)
            decoder = codecs.getincrementaldecoder(name)()
            try:
                while chunk := fh.read(CHUNK):
                    decoder.decode(chunk)
                decoder.decode(b"", final=True)
                return name
            except UnicodeDecodeError:
                continue
    return "latin-1"


def _iter_delimited(info: FileInfo) -> Iterator[list]:
    delimiter = "\t" if info.path.suffix.lower() == ".tsv" else None
    with info.path.open(
        "r", encoding=info.encoding, errors="replace", newline=""
    ) as fh:
        if delimiter is None:
            sample = fh.read(64 * 1024)
            fh.seek(0)
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = ","
        for row in csv.reader(fh, delimiter=delimiter):
            if any(cell.strip() for cell in row):
                yield row


def _iter_xlsx(info: FileInfo) -> Iterator[list]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(
            415, "xlsx support needs openpyxl (pip install openpyxl)"
        ) from exc
    wb = load_workbook(info.path, read_only=True, data_only=True)
    try:
        for row in wb.worksheets[0].iter_rows(values_only=True):
            if any(v is not None and str(v).strip() for v in row):
                yield ["" if v is None else v for v in row]
    finally:
        wb.close()
